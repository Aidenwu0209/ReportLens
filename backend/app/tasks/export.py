"""T12 – Export task (python-docx / PDF).

Renders a structured Word report from the dashboard result and optionally
converts to PDF via LibreOffice CLI.
"""

from __future__ import annotations

import logging
import subprocess
from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core.celery_app import celery_app
from app.core.database import SessionLocal
from app.models.models import (
    EvidenceRef,
    ExportTask,
    LlmResult,
    Metric,
    MetricValue,
    Risk,
)
from app.services.storage import LocalFSAdapter, StoragePaths
from app.utils.ulid import generate_ulid

logger = logging.getLogger(__name__)


def _build_docx(version_id: str, tenant_id: str, db, storage: LocalFSAdapter) -> str:
    """Build a .docx report and return the relative storage path."""
    doc = DocxDocument()

    # Title
    title = doc.add_heading("年报智能分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Summary
    llm_result = db.query(LlmResult).filter_by(version_id=version_id).first()
    if llm_result and llm_result.summary:
        doc.add_heading("一、总体概述", level=1)
        doc.add_paragraph(llm_result.summary)

    # Metrics table
    metric_values = (
        db.query(MetricValue)
        .filter_by(version_id=version_id)
        .all()
    )
    if metric_values:
        doc.add_heading("二、核心指标", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "指标"
        hdr[1].text = "数值"
        hdr[2].text = "单位"
        hdr[3].text = "状态"

        for mv in metric_values:
            metric = db.query(Metric).filter_by(metric_id=mv.metric_id).first()
            row = table.add_row().cells
            row[0].text = metric.name_cn if metric else mv.metric_id
            row[1].text = str(mv.value) if mv.value is not None else (mv.value_text or "-")
            row[2].text = metric.unit if metric and metric.unit else ""
            row[3].text = "需复核" if mv.needs_review else "已确认"

    # Risks
    risks = db.query(Risk).filter_by(version_id=version_id).all()
    if risks:
        doc.add_heading("三、风险提示", level=1)
        for r in risks:
            level_label = {"high": "高", "medium": "中", "low": "低"}.get(r.level, r.level)
            p = doc.add_paragraph()
            run = p.add_run(f"【{level_label}风险】{r.title}")
            run.bold = True
            if r.description:
                doc.add_paragraph(r.description)
            if r.recommendation:
                doc.add_paragraph(f"建议: {r.recommendation}")

            # Evidence
            refs = (
                db.query(EvidenceRef)
                .filter_by(source_type="risk", source_id=r.risk_id)
                .all()
            )
            if refs:
                doc.add_paragraph(
                    "来源: " + ", ".join(f"第{ref.page_no}页" for ref in refs)
                )

    # Recommendations
    if llm_result and llm_result.recommendations:
        doc.add_heading("四、建议", level=1)
        for rec in llm_result.recommendations:
            doc.add_paragraph(rec, style="List Bullet")

    # Save
    export_id = generate_ulid("exp")
    rel_path = StoragePaths.export_file(export_id, "docx")
    abs_path = storage.abs_path(rel_path)
    Path(abs_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(abs_path)

    return rel_path


@celery_app.task(name="app.tasks.export.render_docx_pdf", bind=True, max_retries=2)
def render_docx_pdf(self, export_id: str) -> dict:
    """Render the export file (docx; optionally convert to PDF)."""
    db = SessionLocal()
    storage = LocalFSAdapter()

    try:
        export = db.query(ExportTask).filter_by(export_id=export_id).first()
        if not export:
            raise ValueError(f"ExportTask {export_id} not found")

        export.status = "processing"
        db.commit()

        # Build docx
        rel_path = _build_docx(export.version_id, export.tenant_id, db, storage)
        export.file_path = rel_path

        # Optionally convert to PDF
        if export.format == "pdf":
            abs_docx = storage.abs_path(rel_path)
            try:
                subprocess.run(
                    ["soffice", "--headless", "--convert-to", "pdf", abs_docx,
                     "--outdir", str(Path(abs_docx).parent)],
                    check=True,
                    timeout=120,
                )
                pdf_rel = rel_path.replace(".docx", ".pdf")
                export.file_path = pdf_rel
            except (subprocess.CalledProcessError, FileNotFoundError):
                logger.warning("LibreOffice not available, keeping docx format")

        export.status = "done"
        db.commit()

        logger.info("Export %s completed: %s", export_id, export.file_path)
        return {"export_id": export_id, "file_path": export.file_path}

    except Exception as exc:
        db.rollback()
        if export:
            export.status = "failed"
            export.error_message = str(exc)
            db.commit()
        logger.exception("Export failed: %s", export_id)
        raise self.retry(exc=exc, countdown=15)
    finally:
        db.close()
